"""
PowerPoint 报告生成模块

使用 python-pptx 生成 Snapmaker U1 产品反馈分析报告。
PowerPoint 原生支持 Unicode，无需额外配置中文字体。
"""

import logging
import re
from collections import Counter
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

logger = logging.getLogger(__name__)

# --- 颜色方案 ---
COLOR_PRIMARY = RGBColor(0x1A, 0x52, 0x76)
COLOR_ACCENT = RGBColor(0x29, 0x80, 0xB9)
COLOR_POSITIVE = RGBColor(0x27, 0xAE, 0x60)
COLOR_NEGATIVE = RGBColor(0xE7, 0x4C, 0x3C)
COLOR_NEUTRAL = RGBColor(0x95, 0xA5, 0xA6)
COLOR_MIXED = RGBColor(0xF3, 0x9C, 0x12)
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_DARK = RGBColor(0x33, 0x33, 0x33)
COLOR_LIGHT_BG = RGBColor(0xF0, 0xF7, 0xFF)


# ============================================================
# Helper functions
# ============================================================

def _add_title_slide(prs, title, subtitle=""):
    """添加标题幻灯片。"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # 背景色
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_PRIMARY

    # 产品名
    txBox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "Snapmaker U1"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE
    p.alignment = PP_ALIGN.CENTER

    # 标题
    txBox = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(8), Inches(1))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.color.rgb = COLOR_WHITE
    p.alignment = PP_ALIGN.CENTER

    # 副标题
    if subtitle:
        txBox = slide.shapes.add_textbox(Inches(1), Inches(4.2), Inches(8), Inches(1))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(14)
        p.font.color.rgb = RGBColor(0xCC, 0xDD, 0xEE)
        p.alignment = PP_ALIGN.CENTER

    return slide


def _add_section_slide(prs, title):
    """添加章节分隔幻灯片。"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_ACCENT

    txBox = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(8), Inches(1.5))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE
    p.alignment = PP_ALIGN.CENTER

    return slide


def _add_content_slide(prs, title, body_text, max_chars=1800):
    """添加内容幻灯片（标题 + 正文）。如果内容过长，自动分页。"""
    slides = []
    text = body_text.strip()
    page = 0

    while text:
        page += 1
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slides.append(slide)

        # 标题
        page_title = title if page == 1 else f"{title}（续{page}）"
        txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = page_title
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = COLOR_PRIMARY

        # 分隔线
        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.95),
            Inches(9), Emu(18000)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = COLOR_ACCENT
        line.line.fill.background()

        # 正文内容
        chunk = text[:max_chars]
        text = text[max_chars:]

        # 尽量在段落边界断开
        if text:
            last_nl = chunk.rfind("\n\n")
            if last_nl > max_chars * 0.5:
                text = chunk[last_nl:] + text
                chunk = chunk[:last_nl]
            else:
                last_nl = chunk.rfind("\n")
                if last_nl > max_chars * 0.6:
                    text = chunk[last_nl:] + text
                    chunk = chunk[:last_nl]

        txBox = slide.shapes.add_textbox(
            Inches(0.5), Inches(1.1), Inches(9), Inches(6.2)
        )
        tf = txBox.text_frame
        tf.word_wrap = True

        for line_text in chunk.split("\n"):
            p = tf.add_paragraph()
            line_text = line_text.strip()

            if line_text.startswith("## "):
                p.text = line_text[3:]
                p.font.size = Pt(14)
                p.font.bold = True
                p.font.color.rgb = COLOR_PRIMARY
                p.space_before = Pt(8)
            elif line_text.startswith("### "):
                p.text = line_text[4:]
                p.font.size = Pt(12)
                p.font.bold = True
                p.font.color.rgb = COLOR_ACCENT
                p.space_before = Pt(6)
            elif line_text.startswith("- "):
                p.text = "  " + line_text
                p.font.size = Pt(9)
                p.font.color.rgb = COLOR_DARK
            elif line_text.startswith("| "):
                p.text = line_text
                p.font.size = Pt(8)
                p.font.color.rgb = COLOR_DARK
            else:
                p.text = line_text
                p.font.size = Pt(9)
                p.font.color.rgb = COLOR_DARK

        # 删除第一个空段落
        if tf.paragraphs[0].text == "":
            tf.paragraphs[0]._p.getparent().remove(tf.paragraphs[0]._p)

    return slides


def _add_metrics_slide(prs, title, metrics):
    """添加数据指标卡片幻灯片。metrics: list of (label, value)"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 标题
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    # 指标卡片
    n = len(metrics)
    card_width = 8.5 / n
    for i, (label, value) in enumerate(metrics):
        left = Inches(0.5 + i * card_width + i * 0.1)
        top = Inches(1.5)
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            left, top, Inches(card_width), Inches(2)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLOR_LIGHT_BG
        shape.line.color.rgb = COLOR_ACCENT

        tf = shape.text_frame
        tf.word_wrap = True

        # 数值
        p = tf.paragraphs[0]
        p.text = str(value)
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = COLOR_PRIMARY
        p.alignment = PP_ALIGN.CENTER
        p.space_after = Pt(4)

        # 标签
        p2 = tf.add_paragraph()
        p2.text = label
        p2.font.size = Pt(10)
        p2.font.color.rgb = COLOR_NEUTRAL
        p2.alignment = PP_ALIGN.CENTER

    return slide


def _add_video_table_slide(prs, videos, llm_results):
    """添加视频索引表格幻灯片。"""
    # 每页最多显示 12 个视频
    page_size = 12
    for page_start in range(0, len(videos), page_size):
        page_videos = videos[page_start:page_start + page_size]
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        page_num = page_start // page_size + 1
        total_pages = (len(videos) + page_size - 1) // page_size
        title = f"视频索引（{page_num}/{total_pages}）"

        txBox = slide.shapes.add_textbox(Inches(0.3), Inches(0.2), Inches(9), Inches(0.5))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = COLOR_PRIMARY

        # 表格
        rows = len(page_videos) + 1
        cols = 5
        table = slide.shapes.add_table(
            rows, cols,
            Inches(0.3), Inches(0.8),
            Inches(9.4), Inches(min(6.5, rows * 0.45))
        ).table

        # 列宽
        table.columns[0].width = Inches(0.3)   # #
        table.columns[1].width = Inches(2.0)   # 频道
        table.columns[2].width = Inches(4.0)   # 标题
        table.columns[3].width = Inches(1.2)   # 观看量
        table.columns[4].width = Inches(1.9)   # 赞助

        # 表头
        headers = ["#", "频道", "标题", "观看量", "赞助状态"]
        for j, header in enumerate(headers):
            cell = table.cell(0, j)
            cell.text = header
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_PRIMARY
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(8)
                p.font.bold = True
                p.font.color.rgb = COLOR_WHITE

        # 数据行
        for i, v in enumerate(page_videos):
            idx = page_start + i + 1
            sp_type = v.get("sponsor_status", {}).get("sponsor_type", "unknown")
            sp_map = {
                "sponsored": "赞助", "review_sample": "样机",
                "self_purchased": "自购", "affiliate_only": "联盟",
                "unknown": "-",
            }
            row_data = [
                str(idx),
                v["channel"][:20],
                v["title"][:50],
                f"{v['view_count']:,}",
                sp_map.get(sp_type, sp_type),
            ]
            for j, val in enumerate(row_data):
                cell = table.cell(i + 1, j)
                cell.text = val
                for p in cell.text_frame.paragraphs:
                    p.font.size = Pt(7)
                    p.font.color.rgb = COLOR_DARK


def _extract_section(text, header_patterns):
    """从 LLM Markdown 文本中按标题提取某个 section。"""
    if not text:
        return ""
    for pattern in header_patterns:
        regex = (
            r"##\s*(?:\d+[\.\s]*)?" + pattern +
            r"[^\n]*\n(.*?)(?=\n##\s|\Z)"
        )
        match = re.search(regex, text, re.DOTALL | re.IGNORECASE)
        if match:
            return match.group(1).strip()
    return ""


# ============================================================
# Main entry point
# ============================================================

def generate_pptx_report(
    top_videos,
    filter_stats,
    df_comments,
    llm_results,
    overall_llm,
    output_path,
):
    """
    生成 PowerPoint 产品分析报告。

    参数:
        top_videos: 视频列表 (list[dict])
        filter_stats: 管道统计 (dict)
        df_comments: 评论 DataFrame
        llm_results: {video_id: {transcript: {}, comments: {}}}
        overall_llm: 整体 LLM 分析 (dict)
        output_path: PPTX 输出路径
    """
    output_path = Path(output_path)
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    overall_text = overall_llm.get("analysis_text", "") if overall_llm else ""
    total_comments = len(df_comments) if df_comments is not None else 0
    total_views = sum(v.get("view_count", 0) for v in top_videos)

    # ===== 1. 封面 =====
    subtitle = (
        f"基于 {len(top_videos)} 个评测视频 & "
        f"{total_comments:,} 条用户评论的深度分析\n"
        f"生成日期: {pd.Timestamp.now().strftime('%Y-%m-%d')}"
    )
    _add_title_slide(prs, "YouTube 用户反馈分析报告", subtitle)

    # ===== 2. 数据概览 =====
    transcript_success = sum(
        1 for v in top_videos
        if llm_results.get(v["video_id"], {})
           .get("transcript", {}).get("status") == "success"
    )
    _add_metrics_slide(prs, "数据概览", [
        ("评测视频", len(top_videos)),
        ("有效评论", f"{total_comments:,}"),
        ("总观看量", f"{total_views:,}"),
        ("字幕成功", transcript_success),
    ])

    # ===== 3. 执行摘要 =====
    exec_summary = _extract_section(overall_text, [
        r"执行摘要", r"Executive Summary", r"摘要",
    ])
    if exec_summary:
        _add_content_slide(prs, "执行摘要", exec_summary)

    # ===== 4. 整体口碑评估 =====
    _add_section_slide(prs, "第一部分：评测视频分析")

    reputation = _extract_section(overall_text, [
        r"整体口碑", r"Overall Reputation", r"口碑评估",
    ])
    if reputation:
        _add_content_slide(prs, "整体口碑评估", reputation)

    # ===== 5. 优点统计 =====
    strengths = _extract_section(overall_text, [
        r"优点", r"亮点", r"Strengths",
    ])
    if strengths:
        _add_content_slide(prs, "优点/亮点统计", strengths)

    # ===== 6. 缺点统计 =====
    issues = _extract_section(overall_text, [
        r"缺点", r"问题", r"Issues",
    ])
    if issues:
        _add_content_slide(prs, "缺点/问题统计", issues)

    # ===== 7. U1 vs Bambu =====
    bambu = _extract_section(overall_text, [
        r"Bambu", r"竞品", r"Competitive",
    ])
    if bambu:
        _add_content_slide(prs, "Snapmaker U1 vs Bambu Lab", bambu)

    # ===== 8. 改进方向 =====
    improvements = _extract_section(overall_text, [
        r"改进", r"Improvement",
    ])
    if improvements:
        _add_content_slide(prs, "核心改进方向", improvements)

    # ===== 9. 一致性与分歧 =====
    consensus = _extract_section(overall_text, [
        r"一致性", r"分歧", r"Consensus", r"Disagreement",
    ])
    if consensus:
        _add_content_slide(prs, "评测者观点一致性与分歧", consensus)

    # ===== 10. 市场机会 =====
    market = _extract_section(overall_text, [
        r"市场", r"机会", r"风险", r"Market",
    ])
    if market:
        _add_content_slide(prs, "市场机会与风险", market)

    # ===== 11. 各视频观点摘要 =====
    _add_section_slide(prs, "各视频观点摘要")

    for v in top_videos:
        vid = v["video_id"]
        lr = llm_results.get(vid, {})
        ta_text = lr.get("transcript", {}).get("analysis_text", "")
        if not ta_text:
            continue

        # 提取核心结论 + Bambu对比
        conclusion = _extract_section(ta_text, [
            r"核心结论", r"Core Conclusion", r"总体评价",
        ])
        bambu_cmp = _extract_section(ta_text, [
            r"Bambu", r"竞品对比",
        ])
        brand_pref = _extract_section(ta_text, [
            r"品牌偏好", r"Brand Preference",
        ])

        body = f"频道: {v['channel']}\n"
        body += f"观看量: {v['view_count']:,}\n\n"

        if conclusion:
            body += f"### 核心结论\n{conclusion[:600]}\n\n"
        if bambu_cmp:
            body += f"### Bambu 对比\n{bambu_cmp[:400]}\n\n"
        if brand_pref:
            body += f"### 品牌偏好\n{brand_pref[:300]}\n"

        if not (conclusion or bambu_cmp or brand_pref):
            body += ta_text[:800]

        slide_title = f"{v['channel'][:25]} - {v['title'][:35]}"
        _add_content_slide(prs, slide_title, body)

    # ===== 12. 评论分析 =====
    _add_section_slide(prs, "第二部分：用户评论分析")

    # 评论情感统计
    if df_comments is not None and len(df_comments) > 0:
        counts = df_comments["sentiment"].value_counts()
        sentiment_text = f"总有效评论: {total_comments} 条\n\n"
        label_map = {"positive": "正面", "negative": "负面",
                     "neutral": "中性", "mixed": "混合"}
        for s in ["positive", "neutral", "negative", "mixed"]:
            c = counts.get(s, 0)
            pct = round(c / total_comments * 100, 1) if total_comments else 0
            sentiment_text += f"- {label_map.get(s, s)}: {c} 条 ({pct}%)\n"
        _add_content_slide(prs, "评论情感分布", sentiment_text)

    # LLM 评论分析汇总
    comment_highlights = []
    for v in top_videos:
        vid = v["video_id"]
        ca = llm_results.get(vid, {}).get("comments", {})
        ca_text = ca.get("analysis_text", "")
        if ca_text and ca.get("status") == "success":
            comment_highlights.append(
                f"### {v['channel'][:20]} - {v['title'][:30]}\n"
                f"{ca_text[:500]}..."
            )

    if comment_highlights:
        for i in range(0, len(comment_highlights), 2):
            batch = "\n\n---\n\n".join(comment_highlights[i:i+2])
            _add_content_slide(
                prs,
                f"评论分析详情（{i//2 + 1}）",
                batch,
            )

    # ===== 13. 视频索引 =====
    _add_section_slide(prs, "附录：视频索引")
    _add_video_table_slide(prs, top_videos, llm_results)

    # ===== 保存 =====
    prs.save(str(output_path))
    size_kb = output_path.stat().st_size / 1024
    logger.info(f"  PowerPoint 报告已生成: {output_path} ({size_kb:.0f} KB)")
    return output_path
