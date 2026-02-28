#!/usr/bin/env python3
"""
Reddit Snapmaker U1 用户反馈分析报告生成器
==========================================
读取 reddit_u1_data.json，进行话题分类、情绪分析、问题归类，
生成 PPT 和 PDF 格式的分析报告。

用法:
    python reddit_report.py           # 使用真实数据
    python reddit_report.py --sample  # 使用示例数据测试
"""

import json
import os
import re
import sys
import textwrap
from collections import Counter, defaultdict
from datetime import datetime

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.font_manager as fm
from matplotlib.backends.backend_pdf import PdfPages

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

from reportlab.lib import colors as rl_colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm, cm
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage,
    PageBreak, KeepTogether
)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# 配置
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_FILE = os.path.join(SCRIPT_DIR, "reddit_output", "reddit_u1_data.json")
OUTPUT_DIR = os.path.join(SCRIPT_DIR, "reddit_output")
CHART_DIR = os.path.join(OUTPUT_DIR, "charts")

CHART_COLORS = [
    "#1a73e8", "#ea4335", "#fbbc04", "#34a853", "#9334e6",
    "#ff6d01", "#46bdc6", "#7baaf7", "#f07b72", "#fcd04f",
]

CN_FONT_PATH = "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"
CN_FONT_NAME = "WenQuanYi Zen Hei"
USE_CN = os.path.exists(CN_FONT_PATH)

# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# 字体初始化
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

def setup_fonts():
    """配置 matplotlib 和 reportlab 中文字体"""
    if USE_CN:
        plt.rcParams["font.sans-serif"] = [CN_FONT_NAME, "DejaVu Sans"]
        plt.rcParams["axes.unicode_minus"] = False
        try:
            pdfmetrics.registerFont(TTFont("WQY", CN_FONT_PATH))
        except Exception:
            pass
    else:
        plt.rcParams["font.sans-serif"] = ["DejaVu Sans"]
        plt.rcParams["axes.unicode_minus"] = False


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# 话题分类
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

TOPIC_TYPES = {
    "problem_error": {
        "label": "问题/报错", "label_en": "Problem/Error", "color": "#ea4335",
        "keywords": [
            r"\berror\b", r"\bfail(ed|ure|s)?\b", r"\bissue\b", r"\bproblem\b",
            r"\bbug\b", r"\bcan'?t\b", r"\bwon'?t\b", r"\bstuck\b", r"\bcrash\b",
            r"\bbroken\b", r"\bdefect\b", r"\bnot working\b", r"\bdoesn'?t\b",
            r"\bjam(med|ming)?\b", r"\bclog(ged)?\b", r"\bleak(ing)?\b",
            r"\bwarp(ed|ing)?\b", r"\bgrind(ing)?\b", r"\bnois[ey]\b",
            r"\boffset\b", r"\bmisalign\b", r"\babort(ed)?\b",
        ],
    },
    "question_howto": {
        "label": "疑问/求助", "label_en": "Question/Help", "color": "#fbbc04",
        "keywords": [
            r"\bhow\b", r"\?$", r"\bcan I\b", r"\bis (it|there)\b",
            r"\bdoes\b", r"\bwhat\b", r"\bwhy\b", r"\bwhere\b",
            r"\bhelp\b", r"\badvice\b", r"\brecommend\b",
            r"\banyone\b", r"\bcompatib\b",
        ],
    },
    "feature_request": {
        "label": "功能需求/建议", "label_en": "Feature Request", "color": "#9334e6",
        "keywords": [
            r"\brequest\b", r"\bwish\b", r"\bplease add\b", r"\bfeature\b",
            r"\bimprovement\b", r"\bshould\b", r"\bwould be nice\b",
            r"\bhope\b", r"\bneed\b.*\bsupport\b",
        ],
    },
    "showcase": {
        "label": "作品展示", "label_en": "Showcase", "color": "#34a853",
        "keywords": [
            r"\bshowcase\b", r"\bmy (first|latest|new)\b", r"\bprinted\b",
            r"\bcheck out\b", r"\bjust (got|finished|completed)\b",
        ],
    },
    "tip_guide": {
        "label": "技巧/教程", "label_en": "Tip/Guide", "color": "#46bdc6",
        "keywords": [
            r"\btip\b", r"\btrick\b", r"\bguide\b", r"\btutorial\b",
            r"\bmod(ification)?\b", r"\bDIY\b", r"\bsolution\b",
            r"\bhow.?to\b", r"\bstep.?by.?step\b",
        ],
    },
    "review": {
        "label": "评测/对比", "label_en": "Review/Comparison", "color": "#ff6d01",
        "keywords": [
            r"\breview\b", r"\bvs\.?\b", r"\bcompar(e|ison)\b",
            r"\bfirst impression\b", r"\bunbox\b", r"\bworth\b", r"\bopinion\b",
        ],
    },
    "discussion": {
        "label": "讨论", "label_en": "Discussion", "color": "#9aa0a6",
        "keywords": [
            r"\bdiscuss\b", r"\bthink\b", r"\bfeel\b", r"\bexperience\b",
        ],
    },
}


def classify_post(post):
    """对单个帖子进行话题分类"""
    title = post.get("title", "").lower()
    text = post.get("text", "").lower()
    flair = (post.get("flair") or "").lower()
    scores = defaultdict(float)

    for type_key, info in TOPIC_TYPES.items():
        for pat in info["keywords"]:
            if re.search(pat, title, re.IGNORECASE):
                scores[type_key] += 2.0
            if re.search(pat, text, re.IGNORECASE):
                scores[type_key] += 1.0

    flair_map = {
        "question": "question_howto", "help": "question_howto",
        "issue": "problem_error", "bug": "problem_error",
        "showcase": "showcase", "tip": "tip_guide", "guide": "tip_guide",
        "review": "review", "discussion": "discussion",
    }
    for kw, tp in flair_map.items():
        if kw in flair:
            scores[tp] += 3.0

    if not scores:
        return ("question_howto", "low") if "?" in title else ("discussion", "low")

    best = max(scores, key=scores.get)
    val = scores[best]
    conf = "high" if val >= 4 else "medium" if val >= 2 else "low"
    return best, conf


def classify_all(posts):
    for p in posts:
        tp, conf = classify_post(p)
        p["topic_type"] = tp
        p["topic_type_label"] = TOPIC_TYPES[tp]["label"]
        p["topic_type_label_en"] = TOPIC_TYPES[tp]["label_en"]
        p["type_confidence"] = conf
    return posts


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# 情绪分析
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

POS_KW = [
    r"\blove\b", r"\bamazing\b", r"\bexcellent\b", r"\bgreat\b",
    r"\bperfect\b", r"\bawesome\b", r"\bhappy\b", r"\bbeautiful\b",
    r"\bsolid\b", r"\bworth\b", r"\bbest\b", r"\bthank\b",
    r"\bimpressed\b", r"\bfantastic\b", r"\bsmooth\b", r"\bflawless\b",
    r"\breliable\b", r"\bincredible\b", r"\brecommend\b", r"\bsatisf",
]
NEG_KW = [
    r"\bissue\b", r"\bproblem\b", r"\bbug\b", r"\bfail\b",
    r"\bdisappoint\b", r"\bfrustrat\b", r"\bnois[ey]\b", r"\bloud\b",
    r"\bcrash\b", r"\berror\b", r"\bdefect\b", r"\bbroken\b",
    r"\bexpensive\b", r"\bregret\b", r"\bterrible\b", r"\bwaste\b",
    r"\bhorrible\b", r"\bawful\b", r"\bworse?\b",
    r"\bcan'?t\b", r"\bwon'?t\b", r"\bdoesn'?t\b", r"\bnot working\b",
    r"\bstrugg\b", r"\bannoy\b",
]


def analyze_sentiment(post):
    texts = [post.get("title", ""), post.get("text", "")]
    for c in post.get("comments", []):
        texts.append(c.get("text", ""))
    combined = " ".join(texts).lower()

    pos = sum(1 for p in POS_KW if re.search(p, combined))
    neg = sum(1 for p in NEG_KW if re.search(p, combined))

    if pos > 0 and neg > 0:
        ratio = pos / (pos + neg)
        if ratio > 0.65:
            return "positive", pos, neg
        elif ratio < 0.35:
            return "negative", pos, neg
        return "mixed", pos, neg
    if pos > 0:
        return "positive", pos, neg
    if neg > 0:
        return "negative", pos, neg
    return "neutral", 0, 0


def analyze_all_sentiments(posts):
    for p in posts:
        s, pos, neg = analyze_sentiment(p)
        p["sentiment"] = s
        p["sentiment_pos"] = pos
        p["sentiment_neg"] = neg
    return posts


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# 问题主题提取
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

ISSUE_THEMES = {
    "tool_change":     ("换头/工具切换", "Tool Change",
        [r"\btool.?chang\b", r"\bswap\b", r"\bsnapswap\b", r"\btool\s*head\b"]),
    "print_quality":   ("打印质量", "Print Quality",
        [r"\bquality\b", r"\blayer\b", r"\bstring(ing)?\b", r"\bblob\b", r"\bsurface\b", r"\baccura\b"]),
    "bed_leveling":    ("调平/热床", "Bed Leveling",
        [r"\blevel(ing)?\b", r"\bbed\b", r"\bz.?offset\b", r"\bmesh\b", r"\bfirst.?layer\b"]),
    "filament":        ("耗材/进料", "Filament",
        [r"\bfilament\b", r"\bfeed\b", r"\bretract\b", r"\bclog\b", r"\bjam\b", r"\bextrud"]),
    "slicer":          ("切片软件", "Slicer",
        [r"\bslicer?\b", r"\bluban\b", r"\bcura\b", r"\borca\b", r"\bprusa.?slicer\b", r"\bgcode\b"]),
    "firmware":        ("固件/软件", "Firmware/Software",
        [r"\bfirmware\b", r"\bupdate\b", r"\bwifi\b", r"\bconnect\b", r"\btouchscreen\b"]),
    "noise_vibration": ("噪音/振动", "Noise/Vibration",
        [r"\bnois[ey]\b", r"\bloud\b", r"\bvibrat\b", r"\bquiet\b", r"\binput.?shap\b"]),
    "reliability":     ("可靠性", "Reliability",
        [r"\breliab\b", r"\bstab(le|ility)\b", r"\bconsisten\b", r"\bfail.*print\b", r"\babort\b"]),
    "hardware":        ("硬件/结构", "Hardware",
        [r"\bhardware\b", r"\bmotor\b", r"\bbelt\b", r"\bframe\b", r"\benclosure\b", r"\bfan\b", r"\bsensor\b"]),
    "multi_color":     ("多色打印", "Multi-color",
        [r"\bmulti.?colou?r\b", r"\bmulti.?material\b", r"\bpurge\b", r"\bwipe\b", r"\btower\b"]),
    "shipping_support":("物流/售后", "Shipping/Support",
        [r"\bshipp?(ing|ed)\b", r"\bsupport\b", r"\bwarrant\b", r"\bRMA\b", r"\breturn\b", r"\brefund\b"]),
    "price_value":     ("价格/性价比", "Price/Value",
        [r"\bprice\b", r"\bcost\b", r"\bexpensive\b", r"\bvalue\b", r"\bworth\b", r"\bbudget\b"]),
}


def extract_themes(post):
    title = post.get("title", "")
    text = post.get("text", "")
    combined = f"{title} {text}".lower()
    themes = []
    for key, (_, _, pats) in ISSUE_THEMES.items():
        for pat in pats:
            if re.search(pat, combined, re.IGNORECASE):
                themes.append(key)
                break
    return themes


def extract_all_themes(posts):
    for p in posts:
        p["themes"] = extract_themes(p)
    return posts


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# 参与度评分
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

def compute_engagement(post):
    score = post.get("score", 0)
    comments = post.get("num_comments", 0)
    ratio = post.get("upvote_ratio", 0.5)
    return round(score + comments * 2 + (ratio - 0.5) * 10, 1)


def compute_all_engagement(posts):
    for p in posts:
        p["engagement"] = compute_engagement(p)
    return posts


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# 示例数据（用于无真实数据时的报告演示）
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

def generate_sample_data():
    """生成示例数据供测试"""
    sample_posts = [
        {"id": "s1", "subreddit": "snapmaker", "title": "Snapmaker U1 tool change failure after firmware update",
         "author": "user1", "text": "After updating to the latest firmware, my U1 fails to complete tool changes. The nozzle gets stuck halfway. Anyone else experiencing this issue? Error code shows on the touchscreen.", "url": "https://reddit.com/r/snapmaker/s1",
         "score": 45, "upvote_ratio": 0.92, "num_comments": 23, "created_utc": "2026-02-15 10:30:00 UTC", "flair": "Issue",
         "comments": [
             {"comment_id": "c1", "text": "Same problem here! Tool change fails every time since the update. Very frustrating.", "author": "user2", "score": 12},
             {"comment_id": "c2", "text": "Try rolling back firmware to v3.2.1, that fixed it for me.", "author": "user3", "score": 8},
             {"comment_id": "c3", "text": "Snapmaker support told me they are working on a hotfix.", "author": "user4", "score": 15},
         ]},
        {"id": "s2", "subreddit": "snapmaker", "title": "My first multi-color print on the U1 - Amazing results!",
         "author": "user5", "text": "Just finished my first multi-color print using all 4 tool heads. The quality is incredible! Smooth surfaces, clean color transitions. Love this machine!", "url": "https://reddit.com/r/snapmaker/s2",
         "score": 89, "upvote_ratio": 0.97, "num_comments": 31, "created_utc": "2026-02-20 14:00:00 UTC", "flair": "Showcase",
         "comments": [
             {"comment_id": "c4", "text": "Looks amazing! What slicer settings did you use?", "author": "user6", "score": 5},
             {"comment_id": "c5", "text": "Beautiful print! The U1 is really impressive for multi-color.", "author": "user7", "score": 7},
         ]},
        {"id": "s3", "subreddit": "snapmaker", "title": "How to set up OrcaSlicer with Snapmaker U1?",
         "author": "user8", "text": "I want to use OrcaSlicer instead of Luban for my U1. Does anyone have a guide or profile settings? How do I configure the tool change gcode?", "url": "https://reddit.com/r/snapmaker/s3",
         "score": 32, "upvote_ratio": 0.95, "num_comments": 18, "created_utc": "2026-02-10 08:15:00 UTC", "flair": "Question",
         "comments": [
             {"comment_id": "c6", "text": "Check the Snapmaker wiki, they have OrcaSlicer profiles now.", "author": "user9", "score": 10},
             {"comment_id": "c7", "text": "I wrote a step-by-step guide, will share the link.", "author": "user10", "score": 6},
         ]},
        {"id": "s4", "subreddit": "snapmaker", "title": "U1 bed leveling issues - mesh is inconsistent",
         "author": "user11", "text": "My U1 bed leveling mesh shows huge variations. First layer adhesion is terrible. I've tried cleaning the bed, adjusting z-offset, but nothing helps. The auto bed leveling seems broken.", "url": "https://reddit.com/r/snapmaker/s4",
         "score": 28, "upvote_ratio": 0.88, "num_comments": 15, "created_utc": "2026-02-12 16:45:00 UTC", "flair": "Issue",
         "comments": [
             {"comment_id": "c8", "text": "Check if your probe is clean. Mine had filament residue causing bad readings.", "author": "user12", "score": 9},
             {"comment_id": "c9", "text": "Same issue. Disappointing for a printer at this price point.", "author": "user13", "score": 4},
         ]},
        {"id": "s5", "subreddit": "3Dprinting", "title": "Snapmaker U1 vs Bambu Lab A1 - honest comparison",
         "author": "user14", "text": "I own both printers and wanted to share my honest thoughts. The U1 excels at multi-color with its tool changer approach. Print quality is comparable. U1 is louder but more reliable for multi-material. Bambu is faster for single color. Both are great machines.", "url": "https://reddit.com/r/3Dprinting/s5",
         "score": 156, "upvote_ratio": 0.91, "num_comments": 67, "created_utc": "2026-02-25 11:00:00 UTC", "flair": "Review",
         "comments": [
             {"comment_id": "c10", "text": "Great comparison! How about noise levels?", "author": "user15", "score": 8},
             {"comment_id": "c11", "text": "The tool changer approach is much better for multi-material than AMS.", "author": "user16", "score": 12},
             {"comment_id": "c12", "text": "Worth the price difference?", "author": "user17", "score": 5},
         ]},
        {"id": "s6", "subreddit": "snapmaker", "title": "U1 filament clog in tool head 3 - keeps jamming",
         "author": "user18", "text": "Tool head 3 on my U1 keeps clogging. I've tried different filaments (PLA, PETG) and it always jams after about 30 minutes. The extruder gears are grinding the filament. Is this a hardware defect?", "url": "https://reddit.com/r/snapmaker/s6",
         "score": 19, "upvote_ratio": 0.85, "num_comments": 11, "created_utc": "2026-02-18 09:30:00 UTC", "flair": "Issue",
         "comments": [
             {"comment_id": "c13", "text": "Sounds like a heat creep issue. Check if the fan on tool head 3 is spinning properly.", "author": "user19", "score": 7},
             {"comment_id": "c14", "text": "Had the same problem. RMA'd and got a replacement tool head.", "author": "user20", "score": 3},
         ]},
        {"id": "s7", "subreddit": "snapmaker", "title": "DIY enclosure mod for Snapmaker U1 - reduces noise significantly",
         "author": "user21", "text": "I built a custom enclosure for my U1 using IKEA Lack tables. Added acoustic foam inside. Noise went from annoying to barely noticeable. Here's my step-by-step guide with photos.", "url": "https://reddit.com/r/snapmaker/s7",
         "score": 73, "upvote_ratio": 0.96, "num_comments": 22, "created_utc": "2026-02-22 20:00:00 UTC", "flair": "Tip",
         "comments": [
             {"comment_id": "c15", "text": "This is awesome! Does the enclosure affect print quality with PLA?", "author": "user22", "score": 4},
             {"comment_id": "c16", "text": "Great mod! Now I can print ABS without fumes too.", "author": "user23", "score": 6},
         ]},
        {"id": "s8", "subreddit": "snapmaker", "title": "Feature request: support for Bambu filament RFID on U1",
         "author": "user24", "text": "It would be amazing if Snapmaker could add support for reading Bambu-style RFID filament spools. Auto-detecting filament type and color would improve the multi-color workflow significantly.", "url": "https://reddit.com/r/snapmaker/s8",
         "score": 41, "upvote_ratio": 0.93, "num_comments": 14, "created_utc": "2026-02-08 13:20:00 UTC", "flair": "Discussion",
         "comments": [
             {"comment_id": "c17", "text": "Yes! This would be a game changer for the U1.", "author": "user25", "score": 11},
             {"comment_id": "c18", "text": "I hope they add this in a firmware update.", "author": "user26", "score": 8},
         ]},
        {"id": "s9", "subreddit": "snapmaker", "title": "WiFi connectivity keeps dropping on U1",
         "author": "user27", "text": "My U1 loses WiFi connection every few hours. I have to restart the printer to reconnect. Very annoying when trying to monitor prints remotely. Firmware is latest version.", "url": "https://reddit.com/r/snapmaker/s9",
         "score": 22, "upvote_ratio": 0.87, "num_comments": 9, "created_utc": "2026-02-14 17:00:00 UTC", "flair": "Issue",
         "comments": [
             {"comment_id": "c19", "text": "Same here. WiFi on the U1 is not reliable at all.", "author": "user28", "score": 6},
             {"comment_id": "c20", "text": "Try using a static IP address, that helped me.", "author": "user29", "score": 4},
         ]},
        {"id": "s10", "subreddit": "snapmaker", "title": "U1 shipping delay - still waiting after 3 weeks",
         "author": "user30", "text": "Ordered my U1 three weeks ago and it still hasn't shipped. Customer support keeps saying it will ship soon. Getting worried about this purchase. Anyone else waiting?", "url": "https://reddit.com/r/snapmaker/s10",
         "score": 15, "upvote_ratio": 0.82, "num_comments": 8, "created_utc": "2026-02-05 11:45:00 UTC", "flair": None,
         "comments": [
             {"comment_id": "c21", "text": "Mine took 4 weeks to ship. Be patient, it's worth the wait.", "author": "user31", "score": 3},
             {"comment_id": "c22", "text": "Terrible customer service. I'm considering a refund.", "author": "user32", "score": 5},
         ]},
        {"id": "s11", "subreddit": "3Dprinting", "title": "Is the Snapmaker U1 worth it for multi-color printing?",
         "author": "user33", "text": "Thinking about getting the U1 for multi-color prints. Is it worth the price? How does it compare to Prusa XL or Bambu X1C with AMS? Concerned about reliability and print quality.", "url": "https://reddit.com/r/3Dprinting/s11",
         "score": 38, "upvote_ratio": 0.90, "num_comments": 25, "created_utc": "2026-02-17 15:30:00 UTC", "flair": "Discussion",
         "comments": [
             {"comment_id": "c23", "text": "I love my U1. Best tool changer in this price range.", "author": "user34", "score": 14},
             {"comment_id": "c24", "text": "Great for multi-color, but expect some reliability issues initially.", "author": "user35", "score": 9},
             {"comment_id": "c25", "text": "Bambu AMS is simpler but tool changer is more versatile.", "author": "user36", "score": 7},
         ]},
        {"id": "s12", "subreddit": "snapmaker", "title": "Loud grinding noise during tool change on U1",
         "author": "user37", "text": "Every time the U1 does a tool change, there's a loud grinding noise. It sounds terrible. Is this normal? The prints come out fine but I'm worried about long-term damage.", "url": "https://reddit.com/r/snapmaker/s12",
         "score": 25, "upvote_ratio": 0.89, "num_comments": 13, "created_utc": "2026-02-19 07:00:00 UTC", "flair": "Issue",
         "comments": [
             {"comment_id": "c26", "text": "Some noise is normal during tool change but grinding is not. Contact support.", "author": "user38", "score": 8},
             {"comment_id": "c27", "text": "Check if the tool heads are properly seated in the dock.", "author": "user39", "score": 5},
         ]},
    ]

    return {
        "source": "reddit",
        "subreddits": ["r/snapmaker", "r/3Dprinting"],
        "extraction_time": datetime.now().strftime("%Y-%m-%dT%H:%M:%S"),
        "total_posts": len(sample_posts),
        "total_comments": sum(len(p.get("comments", [])) for p in sample_posts),
        "posts": sample_posts,
        "note": "SAMPLE DATA for demonstration",
    }


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# 图表生成
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

def _label(cn, en):
    return cn if USE_CN else en


def chart_topic_distribution(posts, filepath):
    """话题分类分布图"""
    counter = Counter(p["topic_type"] for p in posts)
    types_ordered = ["problem_error", "question_howto", "feature_request",
                     "showcase", "tip_guide", "review", "discussion"]
    labels = []
    values = []
    colors = []
    for t in types_ordered:
        if counter.get(t, 0) > 0:
            labels.append(_label(TOPIC_TYPES[t]["label"], TOPIC_TYPES[t]["label_en"]))
            values.append(counter[t])
            colors.append(TOPIC_TYPES[t]["color"])

    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(12, 5))

    # 柱状图
    bars = ax1.barh(labels[::-1], values[::-1], color=colors[::-1], height=0.6)
    ax1.set_xlabel(_label("帖子数", "Post Count"))
    ax1.set_title(_label("话题分类分布", "Topic Distribution"))
    for bar, val in zip(bars, values[::-1]):
        ax1.text(bar.get_width() + 0.3, bar.get_y() + bar.get_height()/2,
                 str(val), va="center", fontsize=10, fontweight="bold")

    # 饼图
    ax2.pie(values, labels=labels, colors=colors, autopct="%1.0f%%",
            startangle=90, textprops={"fontsize": 9})
    ax2.set_title(_label("分类占比", "Distribution %"))

    plt.tight_layout()
    plt.savefig(filepath, dpi=150, bbox_inches="tight")
    plt.close()


def chart_sentiment(posts, filepath):
    """情绪分析图"""
    counter = Counter(p["sentiment"] for p in posts)
    order = ["positive", "negative", "mixed", "neutral"]
    labels_map = {
        "positive": (_label("正面", "Positive"), "#34a853"),
        "negative": (_label("负面", "Negative"), "#ea4335"),
        "mixed":    (_label("混合", "Mixed"), "#fbbc04"),
        "neutral":  (_label("中性", "Neutral"), "#9aa0a6"),
    }
    labels = []
    values = []
    colors = []
    for s in order:
        if counter.get(s, 0) > 0:
            lbl, clr = labels_map[s]
            labels.append(lbl)
            values.append(counter[s])
            colors.append(clr)

    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(12, 5))

    # 饼图
    wedges, texts, autotexts = ax1.pie(
        values, labels=labels, colors=colors, autopct="%1.0f%%",
        startangle=90, textprops={"fontsize": 10})
    ax1.set_title(_label("情绪分布", "Sentiment Distribution"))

    # 情绪得分散点
    scores_pos = [p["sentiment_pos"] for p in posts]
    scores_neg = [p["sentiment_neg"] for p in posts]
    sent_colors = [labels_map.get(p["sentiment"], ("", "#999"))[1] for p in posts]
    ax2.scatter(scores_neg, scores_pos, c=sent_colors, alpha=0.7, s=60, edgecolors="white")
    ax2.set_xlabel(_label("负面关键词数", "Negative Keywords"))
    ax2.set_ylabel(_label("正面关键词数", "Positive Keywords"))
    ax2.set_title(_label("情绪象限图", "Sentiment Quadrant"))
    ax2.axhline(y=2, color="#ccc", linestyle="--", linewidth=0.8)
    ax2.axvline(x=2, color="#ccc", linestyle="--", linewidth=0.8)

    plt.tight_layout()
    plt.savefig(filepath, dpi=150, bbox_inches="tight")
    plt.close()


def chart_themes(posts, filepath):
    """问题主题分布图"""
    theme_counter = Counter()
    for p in posts:
        for t in p.get("themes", []):
            theme_counter[t] += 1

    if not theme_counter:
        fig, ax = plt.subplots(figsize=(10, 5))
        ax.text(0.5, 0.5, _label("无主题数据", "No theme data"),
                ha="center", va="center", fontsize=14)
        ax.set_xlim(0, 1)
        ax.set_ylim(0, 1)
        ax.axis("off")
        plt.savefig(filepath, dpi=150, bbox_inches="tight")
        plt.close()
        return

    top = theme_counter.most_common(10)
    labels = [_label(ISSUE_THEMES[k][0], ISSUE_THEMES[k][1]) for k, _ in top]
    values = [v for _, v in top]
    colors = CHART_COLORS[:len(values)]

    fig, ax = plt.subplots(figsize=(10, 5))
    bars = ax.barh(labels[::-1], values[::-1], color=colors[::-1], height=0.6)
    ax.set_xlabel(_label("帖子数", "Post Count"))
    ax.set_title(_label("用户讨论主题 TOP 10", "Top 10 Discussion Themes"))
    for bar, val in zip(bars, values[::-1]):
        ax.text(bar.get_width() + 0.2, bar.get_y() + bar.get_height()/2,
                str(val), va="center", fontsize=10, fontweight="bold")

    plt.tight_layout()
    plt.savefig(filepath, dpi=150, bbox_inches="tight")
    plt.close()


def chart_engagement(posts, filepath):
    """参与度 Top 10 帖子"""
    sorted_posts = sorted(posts, key=lambda p: p.get("engagement", 0), reverse=True)[:10]
    labels = [textwrap.shorten(p["title"], width=40, placeholder="...") for p in sorted_posts]
    values = [p["engagement"] for p in sorted_posts]
    colors = CHART_COLORS[:len(values)]

    fig, ax = plt.subplots(figsize=(10, 5))
    bars = ax.barh(labels[::-1], values[::-1], color=colors[::-1], height=0.6)
    ax.set_xlabel(_label("参与度评分", "Engagement Score"))
    ax.set_title(_label("热门帖子 TOP 10", "Top 10 Posts by Engagement"))
    for bar, val in zip(bars, values[::-1]):
        ax.text(bar.get_width() + 0.5, bar.get_y() + bar.get_height()/2,
                f"{val:.0f}", va="center", fontsize=9, fontweight="bold")

    plt.tight_layout()
    plt.savefig(filepath, dpi=150, bbox_inches="tight")
    plt.close()


def chart_subreddit(posts, filepath):
    """来源 subreddit 分布"""
    counter = Counter(p.get("subreddit", "unknown") for p in posts)
    labels = [f"r/{k}" for k in counter.keys()]
    values = list(counter.values())
    colors = CHART_COLORS[:len(values)]

    fig, ax = plt.subplots(figsize=(6, 4))
    ax.pie(values, labels=labels, colors=colors, autopct="%1.0f%%",
           startangle=90, textprops={"fontsize": 11})
    ax.set_title(_label("来源分布", "Source Distribution"))
    plt.tight_layout()
    plt.savefig(filepath, dpi=150, bbox_inches="tight")
    plt.close()


def generate_all_charts(posts):
    """生成所有图表"""
    os.makedirs(CHART_DIR, exist_ok=True)
    charts = {}

    chart_funcs = {
        "topic_dist": chart_topic_distribution,
        "sentiment": chart_sentiment,
        "themes": chart_themes,
        "engagement": chart_engagement,
        "subreddit": chart_subreddit,
    }
    for name, func in chart_funcs.items():
        path = os.path.join(CHART_DIR, f"{name}.png")
        func(posts, path)
        charts[name] = path
        print(f"  [图表] {path}")

    return charts


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# PPT 生成
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

COLOR_PRIMARY = RGBColor(0x1A, 0x73, 0xE8)
COLOR_DARK = RGBColor(0x20, 0x21, 0x24)
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_GRAY = RGBColor(0x5F, 0x63, 0x68)
COLOR_RED = RGBColor(0xEA, 0x43, 0x35)
COLOR_GREEN = RGBColor(0x34, 0xA8, 0x53)
COLOR_YELLOW = RGBColor(0xFB, 0xBC, 0x04)


def _add_bg(slide, color=RGBColor(0xF8, 0xF9, 0xFA)):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_title_bar(slide, title_text, subtitle_text=""):
    """添加顶部标题栏"""
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(1.1))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_PRIMARY
    bar.line.fill.background()

    tf = bar.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(24)
    p.font.color.rgb = COLOR_WHITE
    p.font.bold = True
    p.alignment = PP_ALIGN.LEFT
    tf.margin_left = Inches(0.5)
    tf.margin_top = Inches(0.15)

    if subtitle_text:
        p2 = tf.add_paragraph()
        p2.text = subtitle_text
        p2.font.size = Pt(12)
        p2.font.color.rgb = RGBColor(0xBB, 0xDE, 0xFB)
        p2.alignment = PP_ALIGN.LEFT


def _add_text_box(slide, left, top, width, height, text, font_size=12,
                  bold=False, color=COLOR_DARK, alignment=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = alignment
    return txBox


def _add_kpi_box(slide, left, top, number, label, color=COLOR_PRIMARY):
    """添加 KPI 指标卡"""
    box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(2.0), Inches(1.2))
    box.fill.solid()
    box.fill.fore_color.rgb = COLOR_WHITE
    box.line.color.rgb = RGBColor(0xDA, 0xDC, 0xE0)
    box.line.width = Pt(1)

    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_top = Inches(0.1)

    p1 = tf.paragraphs[0]
    p1.text = str(number)
    p1.font.size = Pt(28)
    p1.font.bold = True
    p1.font.color.rgb = color
    p1.alignment = PP_ALIGN.CENTER

    p2 = tf.add_paragraph()
    p2.text = label
    p2.font.size = Pt(10)
    p2.font.color.rgb = COLOR_GRAY
    p2.alignment = PP_ALIGN.CENTER


def generate_pptx(posts, stats, charts, filepath):
    """生成 PPT 报告"""
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)
    blank = prs.slide_layouts[6]

    is_sample = stats.get("is_sample", False)

    # ── Slide 1: 封面 ──
    slide = prs.slides.add_slide(blank)
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(5.625))
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_PRIMARY
    bg.line.fill.background()

    _add_text_box(slide, Inches(0.8), Inches(1.2), Inches(8.4), Inches(1.0),
                  "Reddit 用户反馈分析报告", font_size=36, bold=True, color=COLOR_WHITE,
                  alignment=PP_ALIGN.CENTER)
    _add_text_box(slide, Inches(0.8), Inches(2.2), Inches(8.4), Inches(0.6),
                  "Snapmaker U1 3D Printer", font_size=20, color=RGBColor(0xBB, 0xDE, 0xFB),
                  alignment=PP_ALIGN.CENTER)

    subtitle = f"数据来源: r/snapmaker, r/3Dprinting  |  报告日期: {datetime.now().strftime('%Y-%m-%d')}"
    if is_sample:
        subtitle += "  |  ⚠ 示例数据"
    _add_text_box(slide, Inches(0.8), Inches(3.2), Inches(8.4), Inches(0.5),
                  subtitle, font_size=12, color=RGBColor(0x90, 0xCA, 0xF9),
                  alignment=PP_ALIGN.CENTER)

    # ── Slide 2: 数据概览 ──
    slide = prs.slides.add_slide(blank)
    _add_bg(slide)
    _add_title_bar(slide, "数据概览", "Key Metrics Overview")

    _add_kpi_box(slide, Inches(0.4), Inches(1.4), stats["total_posts"], "帖子总数", COLOR_PRIMARY)
    _add_kpi_box(slide, Inches(2.7), Inches(1.4), stats["total_comments"], "评论总数", COLOR_PRIMARY)
    _add_kpi_box(slide, Inches(5.0), Inches(1.4), stats["avg_score"], "平均得分", COLOR_GREEN)
    _add_kpi_box(slide, Inches(7.3), Inches(1.4), f'{stats["avg_upvote"]:.0%}', "平均点赞率", COLOR_GREEN)

    _add_kpi_box(slide, Inches(0.4), Inches(2.9), stats["problem_count"], "问题帖子", COLOR_RED)
    _add_kpi_box(slide, Inches(2.7), Inches(2.9), stats["question_count"], "求助帖子", COLOR_YELLOW)
    _add_kpi_box(slide, Inches(5.0), Inches(2.9), stats["positive_pct"], "正面情绪占比", COLOR_GREEN)
    _add_kpi_box(slide, Inches(7.3), Inches(2.9), stats["negative_pct"], "负面情绪占比", COLOR_RED)

    date_range = f"数据时间范围: {stats.get('date_range', 'N/A')}"
    _add_text_box(slide, Inches(0.4), Inches(4.4), Inches(9.0), Inches(0.4),
                  date_range, font_size=10, color=COLOR_GRAY)

    # ── Slide 3: 话题分类 ──
    slide = prs.slides.add_slide(blank)
    _add_bg(slide)
    _add_title_bar(slide, "话题分类分布", "Topic Classification Distribution")
    if "topic_dist" in charts:
        slide.shapes.add_picture(
            charts["topic_dist"], Inches(0.3), Inches(1.3), Inches(9.4), Inches(4.0))

    # ── Slide 4: 情绪分析 ──
    slide = prs.slides.add_slide(blank)
    _add_bg(slide)
    _add_title_bar(slide, "情绪分析", "Sentiment Analysis")
    if "sentiment" in charts:
        slide.shapes.add_picture(
            charts["sentiment"], Inches(0.3), Inches(1.3), Inches(9.4), Inches(4.0))

    # ── Slide 5: 讨论主题 ──
    slide = prs.slides.add_slide(blank)
    _add_bg(slide)
    _add_title_bar(slide, "用户讨论主题", "Discussion Themes & Issue Categories")
    if "themes" in charts:
        slide.shapes.add_picture(
            charts["themes"], Inches(0.3), Inches(1.3), Inches(9.4), Inches(4.0))

    # ── Slide 6: 热门帖子 ──
    slide = prs.slides.add_slide(blank)
    _add_bg(slide)
    _add_title_bar(slide, "热门帖子 TOP 10", "Most Engaged Posts")
    if "engagement" in charts:
        slide.shapes.add_picture(
            charts["engagement"], Inches(0.3), Inches(1.3), Inches(9.4), Inches(4.0))

    # ── Slide 7: 问题明细表 ──
    problems = [p for p in posts if p["topic_type"] in ("problem_error", "question_howto")]
    problems.sort(key=lambda x: x.get("engagement", 0), reverse=True)
    top_problems = problems[:8]

    if top_problems:
        slide = prs.slides.add_slide(blank)
        _add_bg(slide)
        _add_title_bar(slide, "问题清单", "Issue List (sorted by engagement)")

        headers = ["标题", "分类", "情绪", "得分", "评论", "主题"]
        col_widths = [Inches(3.2), Inches(1.2), Inches(0.8), Inches(0.7), Inches(0.7), Inches(2.8)]
        rows = len(top_problems) + 1
        table = slide.shapes.add_table(rows, 6, Inches(0.3), Inches(1.3),
                                        Inches(9.4), Inches(0.3 * rows)).table
        for i, w in enumerate(col_widths):
            table.columns[i].width = w

        # 表头
        for i, h in enumerate(headers):
            cell = table.cell(0, i)
            cell.text = h
            for para in cell.text_frame.paragraphs:
                para.font.size = Pt(9)
                para.font.bold = True
                para.font.color.rgb = COLOR_WHITE
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_PRIMARY

        # 数据行
        sentiment_map = {"positive": "正面", "negative": "负面", "mixed": "混合", "neutral": "中性"}
        for r, p in enumerate(top_problems, 1):
            vals = [
                textwrap.shorten(p["title"], width=45, placeholder="..."),
                p.get("topic_type_label", ""),
                sentiment_map.get(p.get("sentiment", ""), ""),
                str(p.get("score", 0)),
                str(p.get("num_comments", 0)),
                ", ".join(ISSUE_THEMES.get(t, ("", t, []))[0] for t in p.get("themes", [])[:3]),
            ]
            for c, val in enumerate(vals):
                cell = table.cell(r, c)
                cell.text = val
                for para in cell.text_frame.paragraphs:
                    para.font.size = Pt(8)
                    para.font.color.rgb = COLOR_DARK
                if r % 2 == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor(0xF0, 0xF0, 0xF0)

    # ── Slide 8: 关键发现 ──
    slide = prs.slides.add_slide(blank)
    _add_bg(slide)
    _add_title_bar(slide, "关键发现与建议", "Key Findings & Recommendations")

    findings = _generate_findings(posts, stats)
    y = Inches(1.4)
    for i, (icon, text) in enumerate(findings):
        _add_text_box(slide, Inches(0.5), y, Inches(9.0), Inches(0.35),
                      f"{icon}  {text}", font_size=11, color=COLOR_DARK)
        y += Inches(0.4)

    # ── Slide 9: 来源分布 ──
    slide = prs.slides.add_slide(blank)
    _add_bg(slide)
    _add_title_bar(slide, "数据来源分布", "Source Distribution")
    if "subreddit" in charts:
        slide.shapes.add_picture(
            charts["subreddit"], Inches(2.5), Inches(1.3), Inches(5.0), Inches(3.5))

    prs.save(filepath)
    print(f"  [PPT] {filepath}")


def _generate_findings(posts, stats):
    """根据数据生成关键发现"""
    findings = []

    # 话题分布
    type_counter = Counter(p["topic_type"] for p in posts)
    top_type = type_counter.most_common(1)
    if top_type:
        t, c = top_type[0]
        findings.append(("📊", f"最常见的帖子类型是「{TOPIC_TYPES[t]['label']}」，占比 {c}/{len(posts)} ({c*100//len(posts)}%)"))

    # 情绪
    sent_counter = Counter(p["sentiment"] for p in posts)
    neg = sent_counter.get("negative", 0)
    pos = sent_counter.get("positive", 0)
    if neg > pos:
        findings.append(("⚠️", f"负面情绪帖子 ({neg}) 多于正面 ({pos})，需关注用户痛点"))
    elif pos > neg:
        findings.append(("✅", f"正面情绪帖子 ({pos}) 多于负面 ({neg})，用户整体满意度较高"))

    # 热门主题
    theme_counter = Counter()
    for p in posts:
        for t in p.get("themes", []):
            theme_counter[t] += 1
    top3 = theme_counter.most_common(3)
    if top3:
        themes_str = "、".join(f"{ISSUE_THEMES[k][0]}({v})" for k, v in top3)
        findings.append(("🔍", f"用户最关注的主题: {themes_str}"))

    # 问题帖
    problems = [p for p in posts if p["topic_type"] == "problem_error"]
    if problems:
        findings.append(("🔴", f"共 {len(problems)} 个问题帖子需要关注和跟进"))
        top_problem = max(problems, key=lambda x: x.get("engagement", 0))
        findings.append(("🔥", f"最热门问题: 「{textwrap.shorten(top_problem['title'], 50, placeholder='...')}」(参与度: {top_problem.get('engagement', 0):.0f})"))

    # 功能请求
    feature_reqs = [p for p in posts if p["topic_type"] == "feature_request"]
    if feature_reqs:
        findings.append(("💡", f"共 {len(feature_reqs)} 个功能需求/建议，建议纳入产品路线图评估"))

    # 建议
    findings.append(("📋", "建议: 针对高频问题建立 FAQ 文档，减少重复求助帖"))
    findings.append(("📋", "建议: 积极回应负面反馈帖子，展示官方关注态度"))

    return findings


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# PDF 生成
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

def generate_pdf(posts, stats, charts, filepath):
    """生成 PDF 报告"""
    doc = SimpleDocTemplate(filepath, pagesize=A4,
                            topMargin=1.5*cm, bottomMargin=1.5*cm,
                            leftMargin=2*cm, rightMargin=2*cm)

    # 字体设置
    font_name = "WQY" if USE_CN else "Helvetica"
    font_name_bold = "WQY" if USE_CN else "Helvetica-Bold"

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        "CNTitle", fontName=font_name_bold if USE_CN else font_name,
        fontSize=22, leading=28, alignment=1, spaceAfter=6*mm,
        textColor=rl_colors.HexColor("#1a73e8")))
    styles.add(ParagraphStyle(
        "CNSubtitle", fontName=font_name, fontSize=11, leading=14,
        alignment=1, spaceAfter=8*mm, textColor=rl_colors.HexColor("#5f6368")))
    styles.add(ParagraphStyle(
        "CNH2", fontName=font_name_bold if USE_CN else font_name,
        fontSize=14, leading=20, spaceBefore=6*mm, spaceAfter=3*mm,
        textColor=rl_colors.HexColor("#1a73e8"),
        borderWidth=0, borderPadding=0,
        borderColor=rl_colors.HexColor("#1a73e8")))
    styles.add(ParagraphStyle(
        "CNBody", fontName=font_name, fontSize=10, leading=14,
        spaceAfter=2*mm, textColor=rl_colors.HexColor("#202124")))
    styles.add(ParagraphStyle(
        "CNSmall", fontName=font_name, fontSize=8, leading=11,
        textColor=rl_colors.HexColor("#5f6368")))

    story = []
    is_sample = stats.get("is_sample", False)

    # ── 封面 ──
    story.append(Spacer(1, 3*cm))
    story.append(Paragraph("Reddit 用户反馈分析报告", styles["CNTitle"]))
    story.append(Paragraph("Snapmaker U1 3D Printer", styles["CNSubtitle"]))
    date_str = f"报告日期: {datetime.now().strftime('%Y-%m-%d')}  |  数据来源: r/snapmaker, r/3Dprinting"
    if is_sample:
        date_str += "  |  ⚠ 示例数据"
    story.append(Paragraph(date_str, styles["CNSmall"]))
    story.append(PageBreak())

    # ── 数据概览 ──
    story.append(Paragraph("1. 数据概览", styles["CNH2"]))
    overview_data = [
        ["指标", "数值", "指标", "数值"],
        ["帖子总数", str(stats["total_posts"]), "评论总数", str(stats["total_comments"])],
        ["平均得分", str(stats["avg_score"]), "平均点赞率", f'{stats["avg_upvote"]:.0%}'],
        ["问题帖子", str(stats["problem_count"]), "求助帖子", str(stats["question_count"])],
        ["正面情绪占比", stats["positive_pct"], "负面情绪占比", stats["negative_pct"]],
        ["时间范围", stats.get("date_range", "N/A"), "", ""],
    ]
    t = Table(overview_data, colWidths=[4.5*cm, 3*cm, 4.5*cm, 3*cm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), rl_colors.HexColor("#1a73e8")),
        ("TEXTCOLOR", (0, 0), (-1, 0), rl_colors.white),
        ("FONTNAME", (0, 0), (-1, 0), font_name),
        ("FONTNAME", (0, 1), (-1, -1), font_name),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ALIGN", (1, 0), (1, -1), "CENTER"),
        ("ALIGN", (3, 0), (3, -1), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.5, rl_colors.HexColor("#dadce0")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [rl_colors.white, rl_colors.HexColor("#f8f9fa")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(t)
    story.append(Spacer(1, 5*mm))

    # ── 图表页 ──
    chart_sections = [
        ("2. 话题分类分布", "topic_dist"),
        ("3. 情绪分析", "sentiment"),
        ("4. 讨论主题分布", "themes"),
        ("5. 热门帖子 TOP 10", "engagement"),
    ]
    for title, chart_key in chart_sections:
        if chart_key in charts:
            story.append(PageBreak())
            story.append(Paragraph(title, styles["CNH2"]))
            img = RLImage(charts[chart_key], width=16*cm, height=7*cm)
            story.append(img)
            story.append(Spacer(1, 3*mm))

    # ── 问题清单 ──
    problems = [p for p in posts if p["topic_type"] in ("problem_error", "question_howto")]
    problems.sort(key=lambda x: x.get("engagement", 0), reverse=True)

    if problems:
        story.append(PageBreak())
        story.append(Paragraph("6. 问题清单", styles["CNH2"]))

        sentiment_map = {"positive": "正面", "negative": "负面", "mixed": "混合", "neutral": "中性"}
        table_data = [["#", "标题", "类型", "情绪", "得分", "评论"]]
        for i, p in enumerate(problems[:15], 1):
            table_data.append([
                str(i),
                textwrap.shorten(p["title"], width=40, placeholder="..."),
                p.get("topic_type_label", ""),
                sentiment_map.get(p.get("sentiment", ""), ""),
                str(p.get("score", 0)),
                str(p.get("num_comments", 0)),
            ])

        t = Table(table_data, colWidths=[0.8*cm, 7*cm, 2.5*cm, 1.5*cm, 1.5*cm, 1.5*cm])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), rl_colors.HexColor("#1a73e8")),
            ("TEXTCOLOR", (0, 0), (-1, 0), rl_colors.white),
            ("FONTNAME", (0, 0), (-1, -1), font_name),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, rl_colors.HexColor("#dadce0")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [rl_colors.white, rl_colors.HexColor("#f8f9fa")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        story.append(t)

    # ── 关键发现 ──
    story.append(PageBreak())
    story.append(Paragraph("7. 关键发现与建议", styles["CNH2"]))
    findings = _generate_findings(posts, stats)
    for icon, text in findings:
        story.append(Paragraph(f"{icon}  {text}", styles["CNBody"]))

    doc.build(story)
    print(f"  [PDF] {filepath}")


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# 统计汇总
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

def compute_stats(posts, data):
    total_posts = len(posts)
    total_comments = sum(len(p.get("comments", [])) for p in posts)
    avg_score = round(sum(p.get("score", 0) for p in posts) / max(total_posts, 1), 1)
    avg_upvote = sum(p.get("upvote_ratio", 0) for p in posts) / max(total_posts, 1)

    sent_counter = Counter(p.get("sentiment", "neutral") for p in posts)
    type_counter = Counter(p.get("topic_type", "discussion") for p in posts)

    pos_pct = f"{sent_counter.get('positive', 0) * 100 // max(total_posts, 1)}%"
    neg_pct = f"{sent_counter.get('negative', 0) * 100 // max(total_posts, 1)}%"

    # 时间范围
    dates = []
    for p in posts:
        ts = p.get("created_utc", "")
        if ts:
            try:
                dt = datetime.strptime(ts.replace(" UTC", ""), "%Y-%m-%d %H:%M:%S")
                dates.append(dt)
            except ValueError:
                pass
    date_range = "N/A"
    if dates:
        date_range = f"{min(dates).strftime('%Y-%m-%d')} ~ {max(dates).strftime('%Y-%m-%d')}"

    return {
        "total_posts": total_posts,
        "total_comments": total_comments,
        "avg_score": avg_score,
        "avg_upvote": avg_upvote,
        "problem_count": type_counter.get("problem_error", 0),
        "question_count": type_counter.get("question_howto", 0),
        "positive_pct": pos_pct,
        "negative_pct": neg_pct,
        "date_range": date_range,
        "is_sample": "note" in data and "SAMPLE" in data.get("note", ""),
    }


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# 分析结果 JSON 导出
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

def export_analysis_json(posts, stats, filepath):
    """导出分析结果 JSON"""
    result = {
        "report_time": datetime.now().isoformat(),
        "stats": stats,
        "posts_analyzed": [
            {
                "id": p.get("id"),
                "title": p.get("title"),
                "subreddit": p.get("subreddit"),
                "topic_type": p.get("topic_type"),
                "topic_type_label": p.get("topic_type_label"),
                "sentiment": p.get("sentiment"),
                "themes": p.get("themes"),
                "engagement": p.get("engagement"),
                "score": p.get("score"),
                "num_comments": p.get("num_comments"),
                "url": p.get("url"),
            }
            for p in posts
        ],
    }
    with open(filepath, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)
    print(f"  [JSON] {filepath}")


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# 主流程
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

def main():
    print("=" * 60)
    print("  Reddit Snapmaker U1 用户反馈分析报告生成器")
    print("=" * 60)

    setup_fonts()
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    # 加载数据
    use_sample = "--sample" in sys.argv
    if use_sample:
        print("\n[模式] 使用示例数据生成报告")
        data = generate_sample_data()
    else:
        data = load_data()

    posts = data.get("posts", [])
    if not posts:
        print("[错误] 无帖子数据，无法生成报告")
        print("请先运行 reddit_scraper.py 或使用 --sample 参数")
        sys.exit(1)

    print(f"\n[数据] 共 {len(posts)} 个帖子")

    # 分析流水线
    print("\n[分析] 话题分类...")
    posts = classify_all(posts)
    print("[分析] 情绪分析...")
    posts = analyze_all_sentiments(posts)
    print("[分析] 主题提取...")
    posts = extract_all_themes(posts)
    print("[分析] 参与度评分...")
    posts = compute_all_engagement(posts)

    # 统计
    stats = compute_stats(posts, data)
    print(f"\n[统计] 帖子: {stats['total_posts']}, 评论: {stats['total_comments']}")
    print(f"[统计] 正面: {stats['positive_pct']}, 负面: {stats['negative_pct']}")

    # 图表
    print("\n[图表] 生成图表...")
    charts = generate_all_charts(posts)

    # 报告
    print("\n[报告] 生成报告文件...")
    pptx_path = os.path.join(OUTPUT_DIR, "reddit_analysis_report.pptx")
    pdf_path = os.path.join(OUTPUT_DIR, "reddit_analysis_report.pdf")
    json_path = os.path.join(OUTPUT_DIR, "reddit_analysis.json")

    generate_pptx(posts, stats, charts, pptx_path)
    generate_pdf(posts, stats, charts, pdf_path)
    export_analysis_json(posts, stats, json_path)

    print("\n" + "=" * 60)
    print("  报告生成完成!")
    print(f"  PPT:  {pptx_path}")
    print(f"  PDF:  {pdf_path}")
    print(f"  JSON: {json_path}")
    print("=" * 60)


if __name__ == "__main__":
    main()
